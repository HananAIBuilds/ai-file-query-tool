"""
AI Multi-Format File Query Tool
--------------------------------
Reads a .txt, .pdf, .docx, .csv, or .xlsx file and lets the user
ask natural language questions about its content, powered by the
Gemini API (with automatic Groq fallback if Gemini is unavailable).

Requires the GOOGLE_API_KEY environment variable to be set.
Optionally set GROQ_API_KEY as a fallback if Gemini's quota is exhausted.
"""

import os
from google import genai
from groq import Groq


def read_file(file_name):
    """Detects file type by extension and extracts its text content."""

    if file_name.endswith(".txt"):
        with open(file_name, "r", encoding="utf-8") as file:
            return file.read()

    elif file_name.endswith(".pdf"):
        from pypdf import PdfReader
        reader = PdfReader(file_name)
        text = ""
        for page in reader.pages:
            text += page.extract_text()
        return text

    elif file_name.endswith(".docx"):
        import docx
        doc = docx.Document(file_name)
        text = ""
        for paragraph in doc.paragraphs:
            text += paragraph.text + "\n"
        return text

    elif file_name.endswith(".csv"):
        import csv
        with open(file_name, mode="r", encoding="utf-8") as file:
            reader = csv.reader(file)
            return list(reader)

    elif file_name.endswith(".xlsx"):
        import openpyxl
        workbook = openpyxl.load_workbook(file_name, read_only=True)
        sheet = workbook.active
        data = list(sheet.iter_rows(values_only=True))
        workbook.close()
        return data

    else:
        print("There is no support available for this file type yet.")
        exit()


def generate_response(prompt_text, google_api_key, groq_api_key=None):
    """Try Gemini first; if it fails (quota limit, network issue, etc.), fall back to Groq."""
    try:
        client = genai.Client(api_key=google_api_key)
        response = client.models.generate_content(
            model="gemini-2.5-flash",
            contents=prompt_text
        )
        return response.text

    except Exception as e:
        print(f"[Gemini failed, falling back to Groq] Error: {e}")

        if not groq_api_key:
            return ("I apologize, the AI service (Gemini) is temporarily unavailable, "
                     "and no backup is configured. Please try again shortly.")

        try:
            groq_client = Groq(api_key=groq_api_key)
            groq_response = groq_client.chat.completions.create(
                model="llama-3.1-8b-instant",
                messages=[{"role": "user", "content": prompt_text}]
            )
            return groq_response.choices[0].message.content

        except Exception as groq_error:
            print(f"[Groq also failed] Error: {groq_error}")
            return "I apologize, our AI service is temporarily unavailable. Please try again in a moment."


def main():
    file_name = input("Enter file name (with extension): ")
    my_data = read_file(file_name)
    print("Data received successfully!")

    google_api_key = os.environ["GOOGLE_API_KEY"]
    groq_api_key = os.environ.get("GROQ_API_KEY")  # optional fallback

    query = input("What do you want to know about this file? ")

    prompt = f"""My data is provided below, carefully review it and answer any questions asked by the user in a 
    well-organized and simple manner, strictly based on that data.

User question: {query}

Data:
{my_data}
"""

    answer_text = generate_response(prompt, google_api_key, groq_api_key)

    print("\nAI Response:")
    print(answer_text)


if __name__ == "__main__":
    main()
